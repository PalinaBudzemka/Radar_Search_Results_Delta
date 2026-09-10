import os
from copy import copy
from pathlib import Path
from datetime import datetime

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


timestamp = datetime.now().strftime("%Y-%m-%d_%H-%M-%S")
output_file = f"delta_results/delta_results_{timestamp}.xlsx"


def find_latest_two_pivot_files():
    files = [
        path for path in Path("search_results").glob("*.xlsx")
        if "RADAR Positions" in path.name and not path.name.startswith("~$")
    ]

    if len(files) < 2:
        raise Exception(
            'Need at least two Excel files with "RADAR Positions" in the filename.'
        )

    files = sorted(files, key=lambda path: path.name, reverse=True)

    latest = files[0]
    previous = files[1]

    return latest, previous


def read_positions(file_path):
    df = pd.read_excel(file_path, sheet_name="Positions")

    if "Position ID" not in df.columns:
        raise Exception(f"'Position ID' column not found in {file_path}")

    df["Position ID Compare"] = df["Position ID"].astype(str).str.strip()

    # Remove rows without Position ID
    df = df[df["Position ID Compare"].notna()]
    df = df[df["Position ID Compare"] != ""]

    # Avoid duplicate Position IDs
    df = df.drop_duplicates(subset=["Position ID Compare"]).copy()

    for column in ["Customer", "Project"]:
        if column in df.columns:
            df[column] = df[column].fillna("").astype(str).str.strip()

    return df


def compare_positions(latest_file, previous_file):
    latest_df = read_positions(latest_file)
    previous_df = read_positions(previous_file)

    latest_df["Project Compare"] = latest_df["Project"].fillna("").astype(str).str.strip()
    previous_df["Project Compare"] = previous_df["Project"].fillna("").astype(str).str.strip()

    latest_ids = set(latest_df["Position ID Compare"])
    previous_ids = set(previous_df["Position ID Compare"])

    latest_projects = set(latest_df["Project Compare"])
    previous_projects = set(previous_df["Project Compare"])

    added_ids = latest_ids - previous_ids
    removed_ids = previous_ids - latest_ids
    unchanged_ids = latest_ids & previous_ids

    added_df = latest_df[latest_df["Position ID Compare"].isin(added_ids)].copy()
    added_df["Headcount status"] = added_df["Project Compare"].apply(
        lambda project: "New Project" if project not in previous_projects else "Ramp-Up"
    )

    removed_df = previous_df[previous_df["Position ID Compare"].isin(removed_ids)].copy()
    removed_df["Headcount status"] = removed_df["Project Compare"].apply(
        lambda project: "Project Closure" if project not in latest_projects else "Ramp-Down"
    )

    unchanged_df = latest_df[latest_df["Position ID Compare"].isin(unchanged_ids)].copy()
    unchanged_df["Headcount status"] = "Unchanged"

    delta_df = pd.concat(
        [added_df, removed_df, unchanged_df],
        ignore_index=True
    )

    delta_df = delta_df.drop(columns=["Position ID Compare", "Project Compare"])

    sort_columns = []
    for column in ["Headcount status", "Customer", "Project", "Position ID"]:
        if column in delta_df.columns:
            sort_columns.append(column)

    if sort_columns:
        delta_df = delta_df.sort_values(by=sort_columns)

    return delta_df, len(added_ids), len(removed_ids), len(unchanged_ids), len(latest_ids), len(previous_ids)


def read_summary(file_path):
    workbook = pd.ExcelFile(file_path)
    sheet_names = workbook.sheet_names
    sheet_name = "Tracker" if "Tracker" in sheet_names else "Summary"
    df = pd.read_excel(file_path, sheet_name=sheet_name)

    if not {"Customer", "Project", "Headcount"}.issubset(df.columns):
        positions_df = pd.read_excel(file_path, sheet_name="Positions")
        positions_df["Customer"] = positions_df["Customer"].fillna("").astype(str).str.strip()
        positions_df["Project"] = positions_df["Project"].fillna("").astype(str).str.strip()

        project_counts = (
            positions_df[positions_df["Project"] != ""]
            .groupby(["Customer", "Project"])
            .size()
            .reset_index(name="Headcount")
        )
        customer_counts = (
            positions_df.groupby("Customer")
            .size()
            .reset_index(name="Headcount")
        )
        customer_counts["Project"] = ""
        df = pd.concat([customer_counts, project_counts], ignore_index=True)

    if "Headcount" not in df.columns and "Count of Position ID" in df.columns:
        df = df.rename(columns={"Count of Position ID": "Headcount"})

    required_columns = ["Customer", "Project", "Headcount"]

    for column in required_columns:
        if column not in df.columns:
            raise Exception(f"'{column}' column not found in {sheet_name} sheet of {file_path}")

    df["Customer"] = df["Customer"].fillna("").astype(str).str.strip()
    df["Project"] = df["Project"].fillna("").astype(str).str.strip()
    df["Headcount"] = pd.to_numeric(
        df["Headcount"],
        errors="coerce"
    ).fillna(0).astype(int)

    # Remove Grand Total row
    df = df[df["Customer"] != "Grand Total"].copy()

    return df


def compare_summary_tabs(latest_file, previous_file):
    latest_summary = read_summary(latest_file)
    previous_summary = read_summary(previous_file)

    # ----------------------------
    # Compare Customer totals
    # ----------------------------

    # Customer total rows are rows where Project is empty
    latest_customers = latest_summary[latest_summary["Project"] == ""].copy()
    previous_customers = previous_summary[previous_summary["Project"] == ""].copy()

    latest_customers = latest_customers[["Customer", "Headcount"]].rename(
        columns={"Headcount": "Latest Count"}
    )

    previous_customers = previous_customers[["Customer", "Headcount"]].rename(
        columns={"Headcount": "Previous Count"}
    )

    customer_compare = previous_customers.merge(
        latest_customers,
        on="Customer",
        how="outer"
    )

    customer_compare["Previous Count"] = customer_compare["Previous Count"].fillna(0).astype(int)
    customer_compare["Latest Count"] = customer_compare["Latest Count"].fillna(0).astype(int)
    customer_compare["Delta"] = customer_compare["Latest Count"] - customer_compare["Previous Count"]

    added_customers = customer_compare[
        (customer_compare["Previous Count"] == 0) &
        (customer_compare["Latest Count"] > 0)
    ].copy()

    added_customers["Change"] = "Added"

    removed_customers = customer_compare[
        (customer_compare["Previous Count"] > 0) &
        (customer_compare["Latest Count"] == 0)
    ].copy()

    removed_customers["Change"] = "Removed"

    added_customers = added_customers[
        ["Customer", "Previous Count", "Latest Count", "Delta", "Change"]
    ].sort_values(by=["Latest Count", "Customer"], ascending=[False, True])

    removed_customers = removed_customers[
        ["Customer", "Previous Count", "Latest Count", "Delta", "Change"]
    ].sort_values(by=["Previous Count", "Customer"], ascending=[False, True])

    # Customers existing in both latest and previous files
    existing_customers = customer_compare[
        (customer_compare["Previous Count"] > 0) &
        (customer_compare["Latest Count"] > 0)
    ]["Customer"].tolist()

    # ----------------------------
    # Compare Customer + Project counts
    # ----------------------------

    # Project rows are rows where Project is not empty
    latest_projects = latest_summary[latest_summary["Project"] != ""].copy()
    previous_projects = previous_summary[previous_summary["Project"] != ""].copy()

    latest_projects = latest_projects[
        ["Customer", "Project", "Headcount"]
    ].rename(columns={"Headcount": "Latest Count"})

    previous_projects = previous_projects[
        ["Customer", "Project", "Headcount"]
    ].rename(columns={"Headcount": "Previous Count"})

    project_compare = previous_projects.merge(
        latest_projects,
        on=["Customer", "Project"],
        how="outer"
    )

    project_compare["Previous Count"] = project_compare["Previous Count"].fillna(0).astype(int)
    project_compare["Latest Count"] = project_compare["Latest Count"].fillna(0).astype(int)
    project_compare["Delta"] = project_compare["Latest Count"] - project_compare["Previous Count"]

    # Important:
    # Exclude projects belonging to fully added or fully removed customers.
    # Keep only customers that existed in both files.
    project_compare = project_compare[
        project_compare["Customer"].isin(existing_customers)
    ].copy()

    decreased_projects = project_compare[project_compare["Delta"] < 0].copy()
    decreased_projects["Change"] = "Decreased"

    increased_projects = project_compare[project_compare["Delta"] > 0].copy()
    increased_projects["Change"] = "Increased"

    decreased_projects = decreased_projects[
        ["Customer", "Project", "Previous Count", "Latest Count", "Delta", "Change"]
    ].sort_values(by=["Delta", "Customer", "Project"], ascending=[True, True, True])

    increased_projects = increased_projects[
        ["Customer", "Project", "Previous Count", "Latest Count", "Delta", "Change"]
    ].sort_values(by=["Delta", "Customer", "Project"], ascending=[False, True, True])

    return added_customers, removed_customers, decreased_projects, increased_projects


def build_project_status_reports(delta_positions_df):
    positions = delta_positions_df[delta_positions_df["Project"] != ""].copy()
    rows = []

    for (customer, project), project_rows in positions.groupby(
        ["Customer", "Project"], dropna=False
    ):
        statuses = set(project_rows["Headcount status"])
        previous_count = len(
            project_rows[~project_rows["Headcount status"].isin(["New Project", "Ramp-Up"])]
        )
        latest_count = len(
            project_rows[~project_rows["Headcount status"].isin(["Ramp-Down", "Project Closure"])]
        )

        if statuses == {"New Project"}:
            status = "New Project"
        elif statuses == {"Project Closure"}:
            status = "Project Closure"
        elif "Ramp-Up" in statuses:
            status = "Ramp-Up"
        elif "Ramp-Down" in statuses:
            status = "Ramp-Down"
        else:
            continue

        rows.append({
            "Customer": customer,
            "Project": project,
            "Previous Count": previous_count,
            "Latest Count": latest_count,
            "Delta": latest_count - previous_count,
            "Headcount Status": status,
        })

    project_status_df = pd.DataFrame(rows)
    if project_status_df.empty:
        empty_columns = [
            "Customer", "Project", "Previous Count", "Latest Count", "Delta",
            "Headcount Status"
        ]
        project_status_df = pd.DataFrame(columns=empty_columns)

    project_status_df = project_status_df.sort_values(
        by=["Customer", "Project"], ascending=[True, True]
    )

    return tuple(
        project_status_df[project_status_df["Headcount Status"] == status].copy()
        for status in ["New Project", "Project Closure", "Ramp-Up", "Ramp-Down"]
    )


def format_workbook(file_path):
    wb = load_workbook(file_path)

    header_fill = PatternFill(start_color="D9EAF7", end_color="D9EAF7", fill_type="solid")
    subtotal_fill = PatternFill(start_color="EAF4E2", end_color="EAF4E2", fill_type="solid")
    grand_total_fill = PatternFill(start_color="D9EAF7", end_color="D9EAF7", fill_type="solid")
    added_fill = PatternFill(start_color="D9EAD3", end_color="D9EAD3", fill_type="solid")
    removed_fill = PatternFill(start_color="F4CCCC", end_color="F4CCCC", fill_type="solid")
    increased_fill = PatternFill(start_color="D9EAD3", end_color="D9EAD3", fill_type="solid")
    decreased_fill = PatternFill(start_color="F4CCCC", end_color="F4CCCC", fill_type="solid")

    hidden_sheets = {
        "Delta Positions",
        "New Projects",
        "Project Closure",
        "Ramp-Up",
        "Ramp-Down",
    }

    for ws in wb.worksheets:
        # Header formatting
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.fill = header_fill

        if ws.title == "Tracker":
            tracker_headers = {cell.value: cell.column for cell in ws[1]}
            project_column = tracker_headers.get("Project")
            customer_column = tracker_headers.get("Customer")
            if project_column and customer_column:
                for row_number in range(2, ws.max_row + 1):
                    customer = ws.cell(row=row_number, column=customer_column).value
                    project = ws.cell(row=row_number, column=project_column).value
                    if customer == "Grand Total":
                        fill = grand_total_fill
                    elif project in (None, ""):
                        fill = subtotal_fill
                    else:
                        fill = None

                    if fill:
                        for cell in ws[row_number]:
                            cell.font = Font(bold=True)
                            cell.fill = fill

        headers = [cell.value for cell in ws[1]]

        # Conditional row coloring
        if "Change" in headers:
            change_col_idx = headers.index("Change") + 1

            for row_number in range(2, ws.max_row + 1):
                change_value = ws.cell(row=row_number, column=change_col_idx).value

                if change_value == "Added":
                    fill = added_fill
                elif change_value == "Removed":
                    fill = removed_fill
                elif change_value == "Increased":
                    fill = increased_fill
                elif change_value == "Decreased":
                    fill = decreased_fill
                else:
                    fill = None

                if fill:
                    for cell in ws[row_number]:
                        cell.fill = fill

        if ws.title in hidden_sheets:
            ws.sheet_state = "hidden"

        # Auto-adjust column widths
        for column_cells in ws.columns:
            max_length = 0
            column_letter = get_column_letter(column_cells[0].column)

            for cell in column_cells:
                if cell.value is not None:
                    max_length = max(max_length, len(str(cell.value)))

            ws.column_dimensions[column_letter].width = max_length + 2

    wb.save(file_path)


def build_delta_summary(delta_positions_df, previous_file):
    current_positions = delta_positions_df[
        ~delta_positions_df["Headcount status"].isin(["Ramp-Down", "Project Closure"])
    ].copy()

    project_counts = (
        current_positions.groupby(["Customer", "Project"], dropna=False)["Position ID"]
        .count()
        .reset_index(name="Headcount")
    )

    previous_summary = read_summary(previous_file)
    previous_counts = previous_summary.set_index(["Customer", "Project"])["Headcount"].to_dict()

    project_counts = project_counts[
        project_counts["Project"].notna() & (project_counts["Project"] != "")
    ].copy()

    def get_headcount_status(row):
        previous_count = previous_counts.get((row["Customer"], row["Project"]), 0)
        current_count = row["Headcount"]

        if current_count > 0 and previous_count == 0:
            return "New Project"
        if current_count > previous_count:
            return "Ramp-Up"
        if current_count < previous_count and current_count > 0:
            return "Ramp-Down"
        if current_count == 0 and previous_count > 0:
            return "Project Closure"
        return "Unchanged"

    project_counts["Headcount Status"] = project_counts.apply(get_headcount_status, axis=1)
    project_counts = project_counts.sort_values(
        by=["Headcount", "Customer", "Project"],
        ascending=[False, True, True]
    )

    return project_counts[["Customer", "Project", "Headcount", "Headcount Status"]].reset_index(drop=True)


def highlight_summary_changes(latest_file, previous_file, target_file):
    wb = load_workbook(target_file)
    if "Tracker" not in wb.sheetnames:
        return

    ws = wb["Tracker"]
    headers = [cell.value for cell in ws[1]]
    if "Headcount Status" not in headers:
        return

    status_column = headers.index("Headcount Status") + 1

    increased_fill = PatternFill(
        start_color="D9EAF7", end_color="D9EAF7", fill_type="solid"
    )
    ramp_up_fill = PatternFill(
        start_color="EADCF8", end_color="EADCF8", fill_type="solid"
    )
    ramp_down_fill = PatternFill(
        start_color="D9D9D9", end_color="D9D9D9", fill_type="solid"
    )
    decreased_fill = PatternFill(
        start_color="F4CCCC", end_color="F4CCCC", fill_type="solid"
    )

    for row_number in range(2, ws.max_row + 1):
        status = ws.cell(row=row_number, column=status_column).value

        if status == "New Project":
            fill = increased_fill
        elif status == "Ramp-Up":
            fill = ramp_up_fill
        elif status == "Ramp-Down":
            fill = ramp_down_fill
        elif status == "Project Closure":
            fill = decreased_fill
        else:
            continue

        for cell in ws[row_number]:
            cell.fill = copy(fill)

    wb.save(target_file)


latest_file, previous_file = find_latest_two_pivot_files()

print(f"Latest file:   {latest_file}")
print(f"Previous file: {previous_file}")

delta_positions_df, added_count, removed_count, unchanged_count, latest_total, previous_total = compare_positions(
    latest_file,
    previous_file
)

new_projects_df, project_closure_df, ramp_up_df, ramp_down_df = build_project_status_reports(
    delta_positions_df
)

run_summary_df = pd.DataFrame({
    "Metric": [
        "Latest file",
        "Previous file",
        "Added positions",
        "Removed positions",
        "Unchanged positions",
        "Total positions in latest",
        "Total positions in previous",
        "New projects",
        "Project closures",
        "Ramp-Up projects",
        "Ramp-Down projects"
    ],
    "Value": [
        os.path.basename(latest_file),
        os.path.basename(previous_file),
        added_count,
        removed_count,
        unchanged_count,
        latest_total,
        previous_total,
        len(new_projects_df),
        len(project_closure_df),
        len(ramp_up_df),
        len(ramp_down_df)
    ]
})

with pd.ExcelWriter(output_file, engine="openpyxl") as writer:
    run_summary_df.to_excel(writer, sheet_name="Run Summary", index=False)
    build_delta_summary(delta_positions_df, previous_file).to_excel(
        writer, sheet_name="Tracker", index=False
    )
    delta_positions_df.to_excel(writer, sheet_name="Delta Positions", index=False)
    new_projects_df.to_excel(writer, sheet_name="New Projects", index=False)
    project_closure_df.to_excel(writer, sheet_name="Project Closure", index=False)
    ramp_up_df.to_excel(writer, sheet_name="Ramp-Up", index=False)
    ramp_down_df.to_excel(writer, sheet_name="Ramp-Down", index=False)

format_workbook(output_file)
highlight_summary_changes(latest_file, previous_file, output_file)

print(f"Created {output_file}")