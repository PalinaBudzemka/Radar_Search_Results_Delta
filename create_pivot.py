import pandas as pd
import os
import sys
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from datetime import datetime

if len(sys.argv) > 1:
    input_file = sys.argv[1]
else:
    input_files = [
        path for path in Path(".").rglob("*.xlsx")
        if "RADAR Positions Dashboard" in path.name
    ]
    if not input_files:
        raise FileNotFoundError(
            'No Excel files found with "RADAR Positions Dashboard" in the filename'
        )

    input_file = max(input_files, key=os.path.getmtime)

timestamp = datetime.now().strftime("%Y-%m-%d_%H-%M")
output_file = f"pivot_results/pivot_output_{timestamp}.xlsx"

# Read data from Positions sheet
df = pd.read_excel(input_file, sheet_name="Positions")

# Count positions by Customer and Project
project_counts = (
    df.groupby(["Customer", "Project"], dropna=False)["Position ID"]
    .count()
    .reset_index(name="Headcount")
)

# Sort tracker rows by Headcount descending, with Grand Total kept at the bottom.
project_rows = project_counts.sort_values(
    by=["Headcount", "Customer", "Project"],
    ascending=[False, True, True]
)

summary_rows = []
for _, project_row in project_rows.iterrows():
    summary_rows.append({
        "Customer": project_row["Customer"],
        "Project": project_row["Project"],
        "Headcount": project_row["Headcount"],
        "Row Type": "Project"
    })

# Add Grand Total row
grand_total = project_counts["Headcount"].sum()

summary_rows.append({
    "Customer": "Grand Total",
    "Project": "",
    "Headcount": grand_total,
    "Row Type": "Grand Total"
})

summary_df = pd.DataFrame(summary_rows)

# Keep Row Type only for formatting, not for final visible output
visible_summary_df = summary_df[["Customer", "Project", "Headcount"]]

# Save original data + summary
with pd.ExcelWriter(output_file, engine="openpyxl") as writer:
    df.to_excel(writer, sheet_name="Original Positions", index=False)
    visible_summary_df.to_excel(writer, sheet_name="Tracker", index=False)

# Formatting
wb = load_workbook(output_file)
ws = wb["Tracker"]
ws.sheet_properties.tabColor = "FF00FF00"

header_fill = PatternFill(start_color="D9EAF7", end_color="D9EAF7", fill_type="solid")
subtotal_fill = PatternFill(start_color="EAF4E2", end_color="EAF4E2", fill_type="solid")
grand_total_fill = PatternFill(start_color="D9EAF7", end_color="D9EAF7", fill_type="solid")

# Header formatting
for cell in ws[1]:
    cell.font = Font(bold=True)
    cell.fill = header_fill

# Format grand total row
for excel_row_number, row_type in enumerate(summary_df["Row Type"], start=2):
    if row_type == "Grand Total":
        for cell in ws[excel_row_number]:
            cell.font = Font(bold=True)
            cell.fill = grand_total_fill

# Auto-adjust column widths
for column_cells in ws.columns:
    max_length = 0
    column_letter = get_column_letter(column_cells[0].column)

    for cell in column_cells:
        if cell.value is not None:
            max_length = max(max_length, len(str(cell.value)))

    ws.column_dimensions[column_letter].width = max_length + 2

wb.save(output_file)

print(f"Created {output_file}")